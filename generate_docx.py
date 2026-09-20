import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def add_hyperlink(paragraph, url, text, color="0284C7", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{r_id}"/>')
    new_run = parse_xml(f'<w:r {nsdecls("w")}><w:rPr/><w:t>{text}</w:t></w:r>')
    rPr = new_run.find(qn('w:rPr'))
    if color:
        c = parse_xml(f'<w:color {nsdecls("w")} w:val="{color}"/>')
        rPr.append(c)
    if underline:
        u = parse_xml(f'<w:u {nsdecls("w")} w:val="single"/>')
        rPr.append(u)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def set_cell_margins(cell, top=50, bottom=50, left=0, right=0):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_docx(out_path, photo_path):
    doc = Document()
    
    # 0.5 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        
    c_primary = RGBColor(15, 23, 42)    # Slate 900
    c_accent = RGBColor(2, 132, 199)    # Blue 600
    c_body = RGBColor(30, 41, 59)       # Slate 800
    c_muted = RGBColor(71, 85, 105)     # Slate 600

    # Table Header for Name & Photo
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(6.2), Inches(1.3)]
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width
            
    cell_left = table.cell(0, 0)
    cell_right = table.cell(0, 1)
    cell_right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    
    set_cell_margins(cell_left, top=0, bottom=0, left=0, right=50)
    set_cell_margins(cell_right, top=0, bottom=0, left=50, right=0)
    
    # Name
    p_name = cell_left.paragraphs[0]
    p_name.paragraph_format.space_after = Pt(2)
    p_name.paragraph_format.space_before = Pt(0)
    run_name = p_name.add_run("RIO DARMA FREDIKA, S.Kom.")
    run_name.font.name = "Calibri"
    run_name.font.size = Pt(19)
    run_name.font.bold = True
    run_name.font.color.rgb = c_primary
    
    # Role
    p_role = cell_left.add_paragraph()
    p_role.paragraph_format.space_after = Pt(4)
    p_role.paragraph_format.space_before = Pt(0)
    run_role = p_role.add_run("SOFTWARE & DATA ENGINEER | PENGAWAS OPERASIONAL PERTAMA (POP BNSP)")
    run_role.font.name = "Calibri"
    run_role.font.size = Pt(10)
    run_role.font.bold = True
    run_role.font.color.rgb = c_accent
    
    # Contacts
    p_contact = cell_left.add_paragraph()
    p_contact.paragraph_format.space_after = Pt(0)
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.line_spacing = 1.15
    
    def add_meta(p, label, val_run_builder):
        r_lbl = p.add_run(label + " ")
        r_lbl.font.name = "Calibri"
        r_lbl.font.size = Pt(8.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = c_muted
        val_run_builder(p)
        
    add_meta(p_contact, "Lokasi:", lambda p: p.add_run("Garut, Jawa Barat, Indonesia  |  "))
    add_meta(p_contact, "Email:", lambda p: add_hyperlink(p, "mailto:riodarma789@gmail.com", "riodarma789@gmail.com"))
    
    p_c2 = cell_left.add_paragraph()
    p_c2.paragraph_format.space_after = Pt(0)
    p_c2.paragraph_format.space_before = Pt(0)
    p_c2.paragraph_format.line_spacing = 1.15
    add_meta(p_c2, "Telepon / WhatsApp:", lambda p: p.add_run("+62 812-1423-6050  /  +62 821-1669-8032"))
    
    p_c3 = cell_left.add_paragraph()
    p_c3.paragraph_format.space_after = Pt(0)
    p_c3.paragraph_format.space_before = Pt(0)
    p_c3.paragraph_format.line_spacing = 1.15
    add_meta(p_c3, "Website Portofolio:", lambda p: add_hyperlink(p, "https://riodarma002.github.io/web_porto_rio_darma/", "https://riodarma002.github.io/web_porto_rio_darma/"))
    
    p_c4 = cell_left.add_paragraph()
    p_c4.paragraph_format.space_after = Pt(0)
    p_c4.paragraph_format.space_before = Pt(0)
    p_c4.paragraph_format.line_spacing = 1.15
    add_meta(p_c4, "LinkedIn:", lambda p: add_hyperlink(p, "https://www.linkedin.com/in/rio-darma-fredika-296b4277", "linkedin.com/in/rio-darma-fredika-296b4277"))
    p_c4.add_run("  |  ")
    add_meta(p_c4, "GitHub:", lambda p: add_hyperlink(p, "https://github.com/Riodarma002", "github.com/Riodarma002"))
    p_c4.add_run("  |  ")
    add_meta(p_c4, "Drive:", lambda p: add_hyperlink(p, "https://drive.google.com/drive/folders/1k4P6b37Kgh7yJ0cuBZB62rQASZ9kA5oE?usp=sharing", "Folder Portofolio"))
    
    # Photo in right cell
    p_pic = cell_right.paragraphs[0]
    p_pic.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if os.path.exists(photo_path):
        p_pic.add_run().add_picture(photo_path, width=Inches(1.05))
        
    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title)
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = c_primary
        
        # Border bottom via XML
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="2" w:color="0284C7"/></w:pBdr>')
        pPr.append(pBdr)

    # 1. RINGKASAN PROFESIONAL
    add_section_header("RINGKASAN PROFESIONAL")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_before = Pt(3)
    p_sum.paragraph_format.space_after = Pt(6)
    p_sum.paragraph_format.line_spacing = 1.15
    run_sum = p_sum.add_run(
        "Software & Data Engineer dengan pengalaman profesional lebih dari 7 tahun dalam "
        "merancang aplikasi desktop dan platform web operasional, arsitektur data terpadu, otomatisasi pipeline ETL, "
        "serta sistem informasi geospasial (GIS). Berpengalaman memimpin inisiatif digitalisasi operasional armada di "
        "PT. Mega Global Energy melalui perancangan Fleet Management System (FMS), pelacakan GPS armada terintegrasi, "
        "modul geofencing, serta visualisasi data analitik berbasis Python, Vue.js, dan Power BI. "
        "Memiliki rekam jejak enterprise dalam pemodelan data geospasial kelistrikan ArcGIS untuk sinkronisasi aset ke "
        "IBM MAXIMO PT PLN (Persero) bersama PT. Quadran Inovasi Karya Bersama, riset antarmuka pengguna (UX Research) di "
        "PT. Telkom Indonesia, serta mengantongi sertifikasi kompetensi resmi Pengawas Operasional Pertama (POP) dari BNSP. "
        "Terbukti andal dalam menghubungkan kebutuhan operasional lapangan dengan arsitektur teknologi yang efisien, terukur, dan akurat."
    )
    run_sum.font.name = "Calibri"
    run_sum.font.size = Pt(8.5)
    run_sum.font.color.rgb = c_body

    # 2. KOMPETENSI INTI
    add_section_header("KOMPETENSI INTI & KEAHLIAN TEKNIS")
    
    skills = [
        ("Software & Web Engineering:", "Python, JavaScript, Vue.js (Vue 3, Vite, Pinia), HTML5, CSS3, RESTful APIs, Desktop App Development (Tkinter, PyQt), Git & GitHub Actions CI/CD."),
        ("Data Engineering & GIS:", "ArcGIS (Spatial Network, Digitasi Jaringan Listrik & Kontur), MySQL, PostgreSQL, ETL Automation Pipeline, Data Modeling, IBM MAXIMO Enterprise Integration, Microsoft Power BI."),
        ("AI Tooling & Otomasi:", "Antigravity Agentic Workflows, Codex, Ollama (Local LLM), Claude Code, Hermes Agent, 9 Router, Prompt & Pipeline Engineering."),
        ("UI/UX & Desain Produk:", "Figma UI/UX Design, Heuristic Evaluation, User Flow & Information Architecture, Usability Testing, Adobe Photoshop, Adobe Illustrator."),
        ("Operasional & Manajemen:", "Sertifikasi BNSP POP (Pengawas Operasional Pertama), ERP Odoo, Advanced Microsoft Excel (Macro & Formulas), Google Sheets, Fleet Management Systems (FMS), Tata Kelola Kontrak Proyek & LPJ.")
    ]
    
    for cat, val in skills:
        p_sk = doc.add_paragraph()
        p_sk.paragraph_format.space_before = Pt(1)
        p_sk.paragraph_format.space_after = Pt(1)
        p_sk.paragraph_format.line_spacing = 1.15
        
        r1 = p_sk.add_run(cat + " ")
        r1.font.name = "Calibri"
        r1.font.size = Pt(8.5)
        r1.font.bold = True
        r1.font.color.rgb = c_primary
        
        r2 = p_sk.add_run(val)
        r2.font.name = "Calibri"
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = c_body

    # 3. PENGALAMAN KERJA
    add_section_header("PENGALAMAN KERJA")
    
    def add_job(role, company, period, loc, bullets):
        t = doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        t.rows[0].cells[0].width = Inches(5.5)
        t.rows[0].cells[1].width = Inches(2.0)
        
        c1 = t.cell(0, 0)
        c2 = t.cell(0, 1)
        set_cell_margins(c1, top=30, bottom=10, left=0, right=0)
        set_cell_margins(c2, top=30, bottom=10, left=0, right=0)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.keep_with_next = True
        r_role = p1.add_run(role + "\n")
        r_role.font.name = "Calibri"
        r_role.font.size = Pt(9.5)
        r_role.font.bold = True
        r_role.font.color.rgb = c_primary
        
        r_comp = p1.add_run(company)
        r_comp.font.name = "Calibri"
        r_comp.font.size = Pt(9)
        r_comp.font.bold = True
        r_comp.font.color.rgb = c_primary
        
        if loc:
            r_loc = p1.add_run(f" | {loc}")
            r_loc.font.name = "Calibri"
            r_loc.font.size = Pt(8.5)
            r_loc.font.italic = True
            r_loc.font.color.rgb = c_muted
            
        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.keep_with_next = True
        r_per = p2.add_run(period)
        r_per.font.name = "Calibri"
        r_per.font.size = Pt(8.5)
        r_per.font.bold = True
        r_per.font.color.rgb = c_muted

        for b in bullets:
            p_b = doc.add_paragraph(style='List Bullet')
            p_b.paragraph_format.space_before = Pt(1)
            p_b.paragraph_format.space_after = Pt(2)
            p_b.paragraph_format.line_spacing = 1.15
            run_b = p_b.add_run(b)
            run_b.font.name = "Calibri"
            run_b.font.size = Pt(8.5)
            run_b.font.color.rgb = c_body

    # Job 1
    add_job(
        "Software & Data Engineer (Foreman Planning)",
        "PT. MEGA GLOBAL ENERGY",
        "Desember 2023 – Sekarang",
        "PT. GAM & PT. INDEXIM / Site Operation",
        [
            "Merancang, membangun, dan memelihara aplikasi desktop serta platform web operasional internal (Fleet Management System/FMS, Monitoring Bus, Monitoring Dump Truck, dan Monitoring Fueltruck) untuk pemantauan unit secara real-time.",
            "Mengembangkan arsitektur pelacakan GPS armada terintegrasi dan modul geofencing virtual guna membatasi serta melacak mobilitas unit operasional tambang dengan presisi tinggi.",
            "Membangun dashboard Business Intelligence interaktif menggunakan Power BI, Python, dan Vue.js untuk analisis capaian target produksi harian, status armada, serta efisiensi jam kerja operasional.",
            "Mengembangkan platform web Daily MS Rental (planning.mge.co.id) dan aplikasi dispatch untuk mengawasi kondisi unit running per-shift, kendala delay, dan breakdown armada.",
            "Melakukan pemrosesan dan analisis data harian konsumsi bahan bakar (fuel consumption) dan produktivitas fleet untuk mendukung pengambilan keputusan strategis manajemen operasional."
        ]
    )

    # Job 2
    add_job(
        "Data Engineer & Drafter GIS Jaringan Listrik",
        "PT. QUADRAN INOVASI KARYA BERSAMA",
        "September 2021 – Oktober 2023",
        "Bandung, Jawa Barat (Mitra PT PLN)",
        [
            "Mengolah, memverifikasi, dan memodelkan data survei koordinat lapangan ke dalam sistem informasi geografis ArcGIS untuk pemetaan jaringan distribusi kelistrikan berskala enterprise.",
            "Memetakan dan mendigitasi komponen jaringan kelistrikan tegangan menengah dan rendah (Gardu Induk, SUTM, JTR, Gardu Distribusi CBTG U, Switch) bermitra langsung dengan PT PLN (Persero).",
            "Mengembangkan skrip otomasi pipeline konversi data spasial (AR text format) menjadi data tabular terstruktur (Excel/CSV) untuk sinkronisasi aset ke sistem enterprise IBM MAXIMO PLN.",
            "Melakukan validasi dan verifikasi teknis kesesuaian gambar as-built dengan konstruksi fisik aktual di lapangan guna menjamin integritas database aset kelistrikan."
        ]
    )

    # Job 3
    add_job(
        "Administrasi Umum & Pengelola LPJ",
        "PT. KA PROPERTI MANAJEMEN",
        "September 2019 – Mei 2021",
        "Proyek Reaktivasi Cibatu – Garut",
        [
            "Mengelola tertib administrasi berkas legal konstruksi proyek reaktivasi jalur kereta api, meliputi SPK, BAP, BAST basborong, dan surat perjanjian kontrak pegawai.",
            "Menyusun Laporan Pertanggungjawaban (LPJ) keuangan dan verifikasi progres fisik pekerjaan borongan sesuai jadwal dan standar audit PT KA Properti Manajemen.",
            "Memastikan kelancaran alur verifikasi berkas dan kepatuhan administratif antara vendor kontraktor pelaksana, konsultan, dan manajemen pusat."
        ]
    )

    # Job 4
    add_job(
        "Admin & Pengawas Alat Berat",
        "PT. SUMBER DAYA ARGA (SDA)",
        "Juni 2019 – Desember 2019",
        "Area Proyek Lapangan",
        [
            "Bertanggung jawab atas pengawasan operasional harian, keselamatan kerja, dan progres pergerakan unit alat berat di area proyek konstruksi.",
            "Mengelola time sheet kerja operator, log konsumsi bahan bakar solar, serta rekapitulasi penggajian berbasis jam kerja operasional unit.",
            "Menyusun laporan harian kondisi mesin dan jadwal pemeliharaan berkala guna meminimalkan downtime alat di lapangan."
        ]
    )

    # Job 5
    add_job(
        "Staff Research & Development (R&D)",
        "PT. MACO GROUP",
        "Januari 2019 – Agustus 2019",
        "Bandung, Jawa Barat",
        [
            "Merancang prototipe antarmuka UI/UX website e-commerce Catenzo dan kanal digital produk Maco Group berdasarkan prinsip user experience dan kemudahan navigasi.",
            "Membuat konsep visual desain produk kreatif, materi promosi digital, dan visual publikasi media sosial."
        ]
    )

    # Job 6
    add_job(
        "UX Researcher & Graphic Designer / CRM Service Desk",
        "PT. TELKOM INDONESIA / PT. INFOMEDIA SOLUSI HUMANIKA",
        "April 2017 – Desember 2018",
        "Bandung, Jawa Barat",
        [
            "Menganalisis antarmuka produk digital (web & mobile) internal Telkom dengan metode heuristic evaluation dan usability testing.",
            "Menyusun alur pengguna (user flow), arsitektur informasi (IA), dan wireframe prototipe untuk kolaborasi bersama tim developer.",
            "Mengolah hasil masukan pengujian pengguna (user feedback) menjadi dokumen rekomendasi perbaikan desain antarmuka.",
            "Menangani tiket eskalasi dan layanan CRM pada unit HCC DDS ISH Bandung dengan standar penyelesaian tepat waktu."
        ]
    )

    # 4. PENDIDIKAN FORMAL
    add_section_header("PENDIDIKAN FORMAL")
    
    t_edu = doc.add_table(rows=2, cols=2)
    t_edu.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_edu.autofit = False
    t_edu.rows[0].cells[0].width = Inches(5.8)
    t_edu.rows[0].cells[1].width = Inches(1.7)
    t_edu.rows[1].cells[0].width = Inches(5.8)
    t_edu.rows[1].cells[1].width = Inches(1.7)
    
    # Edu 1
    p_e1 = t_edu.cell(0, 0).paragraphs[0]
    p_e1.paragraph_format.space_after = Pt(2)
    r_ed1 = p_e1.add_run("STMIK \"AMIKBANDUNG\" — Sarjana Komputer (S.Kom.), Teknik Informatika\n")
    r_ed1.font.name = "Calibri"
    r_ed1.font.size = Pt(9)
    r_ed1.font.bold = True
    r_ed1.font.color.rgb = c_primary
    
    r_ed1_sub = p_e1.add_run("Status: Terakreditasi B (BAN-PT) | IPK: 2.79 / 4.00 (145 SKS)\n"
                             "Skripsi: Perancangan dan Implementasi Aplikasi Guitar Guide Sebagai Media Pembelajaran Untuk Grade Pertama Kursus Gitar Berbasis Mobile")
    r_ed1_sub.font.name = "Calibri"
    r_ed1_sub.font.size = Pt(8.5)
    r_ed1_sub.font.color.rgb = c_muted
    r_ed1_sub.font.italic = True
    
    p_e1_d = t_edu.cell(0, 1).paragraphs[0]
    p_e1_d.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_ed1_d = p_e1_d.add_run("2011 – 2015\nBandung, Jawa Barat")
    r_ed1_d.font.name = "Calibri"
    r_ed1_d.font.size = Pt(8.5)
    r_ed1_d.font.bold = True
    r_ed1_d.font.color.rgb = c_muted

    # Edu 2
    p_e2 = t_edu.cell(1, 0).paragraphs[0]
    p_e2.paragraph_format.space_before = Pt(3)
    p_e2.paragraph_format.space_after = Pt(2)
    r_ed2 = p_e2.add_run("SMA NEGERI 3 GARUT — Jurusan Ilmu Pengetahuan Alam (IPA)")
    r_ed2.font.name = "Calibri"
    r_ed2.font.size = Pt(9)
    r_ed2.font.bold = True
    r_ed2.font.color.rgb = c_primary
    
    p_e2_d = t_edu.cell(1, 1).paragraphs[0]
    p_e2_d.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_e2_d.paragraph_format.space_before = Pt(3)
    r_ed2_d = p_e2_d.add_run("2007 – 2010\nGarut, Jawa Barat")
    r_ed2_d.font.name = "Calibri"
    r_ed2_d.font.size = Pt(8.5)
    r_ed2_d.font.bold = True
    r_ed2_d.font.color.rgb = c_muted

    # 5. SERTIFIKASI
    add_section_header("SERTIFIKASI RESMI & PELATIHAN PROFESIONAL")
    certs = [
        "Sertifikasi Kompetensi BNSP — Pengawas Operasional Pertama (POP) | LSP Energi Mandiri | No. Sertifikat: 05100 3121 4 0035925 2025 | Masa Berlaku: Mei 2025 – 26 Mei 2030",
        "Pelatihan Profesional UI/UX & Usability Research — Komunitas UX.ID Bandung (03 Mei 2018)",
        "Pelatihan Character Building & Kepemimpinan — STMIK AMIKBANDUNG (21 Maret 2015)"
    ]
    for c in certs:
        p_c = doc.add_paragraph(style='List Bullet')
        p_c.paragraph_format.space_before = Pt(1)
        p_c.paragraph_format.space_after = Pt(2)
        r = p_c.add_run(c)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.color.rgb = c_body

    # 6. PORTOFOLIO
    add_section_header("PORTOFOLIO TERPILIH & VERIFIKASI KARYA")
    p_po = doc.add_paragraph()
    p_po.paragraph_format.space_before = Pt(2)
    p_po.paragraph_format.space_after = Pt(2)
    r_po = p_po.add_run("Portofolio lengkap dan visualisasi sistem operasional dapat diakses secara interaktif pada website portofolio: ")
    r_po.font.name = "Calibri"
    r_po.font.size = Pt(8.5)
    r_po.font.color.rgb = c_body
    add_hyperlink(p_po, "https://riodarma002.github.io/web_porto_rio_darma/", "https://riodarma002.github.io/web_porto_rio_darma/")

    portos = [
        "Fleet Management System (FMS) & Geofencing: Solusi desktop dan virtual geofence untuk pelacakan koordinat dan efisiensi mobilitas armada secara real-time.",
        "Dashboard Web Optrack & Produksi: Sistem visualisasi performa armada dan evaluasi target produksi tambang (planning.mge.co.id).",
        "Digitalisasi Spasial ArcGIS & MAXIMO PLN: Pemodelan jaringan distribusi Gardu CBTG U, SUTM, JTR, dan konversi data tabular ke enterprise MAXIMO.",
        "Dashboard Business Intelligence Power BI: Pemantauan unit commissioning, distribusi logistik bahan bakar Fueltruck, dan efisiensi operasional."
    ]
    for po in portos:
        p_poi = doc.add_paragraph(style='List Bullet')
        p_poi.paragraph_format.space_before = Pt(1)
        p_poi.paragraph_format.space_after = Pt(2)
        r = p_poi.add_run(po)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.color.rgb = c_body

    # 7. REFERENSI
    add_section_header("REFERENSI PROFESIONAL")
    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.space_before = Pt(2)
    p_ref.paragraph_format.space_after = Pt(4)
    r_rf1 = p_ref.add_run("Gilan Khusnul — Drafter GIS, PT. Quadran Inovasi Karya Bersama | Telepon: +62 859-4670-6123 | Email: ")
    r_rf1.font.name = "Calibri"
    r_rf1.font.size = Pt(8.5)
    r_rf1.font.color.rgb = c_body
    add_hyperlink(p_ref, "mailto:gilankhusnul022@gmail.com", "gilankhusnul022@gmail.com")

    doc.save(out_path)
    print(f"Successfully generated DOCX: {out_path}")

if __name__ == '__main__':
    photo = 'dokumenrio/foto_rio_formal.jpg'
    out_docx = 'dokumenrio/CV_ATS_Rio_Darma_Fredika.docx'
    create_docx(out_docx, photo)
